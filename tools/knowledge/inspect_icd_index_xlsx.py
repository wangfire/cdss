# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

path = r"D:\CDSS智能编码系统\docs\superpowers\konwledge\ICD索引整理_编码规则修正版_最终版.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n=== %s max_row=%s max_col=%s" % (name, ws.max_row, ws.max_column))
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=4, values_only=True), 1):
        print(i, r)
wb.close()
